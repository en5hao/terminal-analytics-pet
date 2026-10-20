#!/usr/bin/env python3
"""Превью-картинки для screenshots/. ВАЖНО: это РЕНДЕРЫ результатов (matplotlib) по данным проекта,
а не снимки экрана Power BI / Power Query / Excel. Каждая картинка подписана этим в подвале.
Перед публикацией замените их реальными скриншотами из приложений (имена файлов те же)."""
import subprocess
from pathlib import Path
import io
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import pandas as pd
from metrics import compute

plt.rcParams["font.family"] = "DejaVu Sans"
ROOT = Path(__file__).resolve().parent.parent
S = ROOT / "screenshots"
k, D = compute()
NAVY, BLUE, RED, GREY, BG = "#1F3A5F", "#2F6DB5", "#C8442F", "#6B7280", "#F4F6F9"
NOTE = "Рендер по данным проекта (matplotlib), не снимок экрана приложения — заменить реальным скриншотом"

def footer(fig, text=NOTE):
    fig.text(0.5, 0.008, text, ha="center", fontsize=7.5, color=GREY, style="italic")

def table(ax, df, fs=8.5, colw=None, hl=None):
    vals = [[str(v).replace("$", r"\$") for v in row] for row in df.values]
    tb = ax.table(cellText=vals, colLabels=list(df.columns), loc="upper center", cellLoc="left")
    tb.auto_set_font_size(False); tb.set_fontsize(fs); tb.scale(1, 1.45)
    if not colw: tb.auto_set_column_width(list(range(len(df.columns))))
    for (r, c), cell in tb.get_celld().items():
        cell.set_edgecolor("#D5DAE1")
        if r == 0: cell.set_facecolor(NAVY); cell.get_text().set_color("white"); cell.get_text().set_weight("bold")
        elif hl and hl(r - 1): cell.set_facecolor("#FBE9E6")
    if colw: 
        for c, w in enumerate(colw):
            for r in range(len(df) + 1): tb[r, c].set_width(w)
    ax.axis("off")

def kpi(ax, label, value, color=NAVY):
    ax.set_facecolor("white"); ax.set_xticks([]); ax.set_yticks([])
    for s in ax.spines.values(): s.set_color("#D5DAE1")
    ax.text(0.06, 0.70, label, transform=ax.transAxes, fontsize=10, color=GREY)
    ax.text(0.06, 0.20, value, transform=ax.transAxes, fontsize=24, color=color, weight="bold")

# ------------------------------------------------------------- Overview
daily = pd.DataFrame(k["daily"]); daily["d"] = pd.to_datetime(daily["txn_time"])
an = D["dc"].query("anomaly_reason != 'норма'")[["device_id", "model", "location", "anomaly_reason"]]
fig = plt.figure(figsize=(16, 9), facecolor=BG)
fig.text(0.03, 0.95, "Terminal Analytics — Overview", fontsize=20, weight="bold", color=NAVY)
for i, (lab, opts) in enumerate([("txn_time", "01.09.2026 — 30.09.2026"), ("model", "Все"), ("location", "Все")]):
    fig.text(0.50 + i * 0.165, 0.955, f"{lab}:  {opts}", fontsize=9.5, color=NAVY,
             bbox=dict(boxstyle="round,pad=0.4", fc="white", ec="#D5DAE1"))
vals = [("Success Rate", f"{k['success_rate']:.1%}"), ("MTTR Hours", f"{k['mttr_hours']:.2f}"),
        ("Anomaly Devices", str(k["anomaly_devices"])), ("Total Txns", str(k["total_txns"]))]
for i, (l, v) in enumerate(vals):
    kpi(fig.add_axes([0.03 + i * 0.2425, 0.78, 0.225, 0.125]), l, v, RED if l == "Anomaly Devices" else NAVY)
ax = fig.add_axes([0.04, 0.42, 0.45, 0.30], facecolor="white")
ax.plot(daily["d"], daily["n"], color=BLUE, lw=2); ax.fill_between(daily["d"], daily["n"], alpha=0.08, color=BLUE)
ax.set_title("Транзакции по дням", loc="left", fontsize=11, color=NAVY); ax.grid(alpha=0.25)
import matplotlib.dates as mdates
ax.xaxis.set_major_locator(mdates.DayLocator(interval=4)); ax.xaxis.set_major_formatter(mdates.DateFormatter("%d.%m")); ax.tick_params(labelsize=8)
for s in ("top", "right"): ax.spines[s].set_visible(False)
top = pd.DataFrame(k["top_error_devices"])
ax = fig.add_axes([0.56, 0.42, 0.40, 0.30], facecolor="white")
ax.barh(top["device_id"][::-1], top["errors"][::-1], color=RED); ax.set_title("Топ-5 устройств по числу ошибок (failed + timeout)", loc="left", fontsize=11, color=NAVY)
for i, v in enumerate(top["errors"][::-1]): ax.text(v + 0.2, i, str(int(v)), va="center", fontsize=9)
for s in ("top", "right"): ax.spines[s].set_visible(False)
ax = fig.add_axes([0.04, 0.06, 0.92, 0.28]); ax.set_title("Аномальные устройства (device_consistency)", loc="left", fontsize=11, color=NAVY)
table(ax, an, fs=9, colw=[0.12, 0.12, 0.30, 0.46])
footer(fig); fig.savefig(S / "powerbi" / "overview.png", dpi=110); plt.close(fig)

# --------------------------------------------------------------- Detail
dev = "ATM-031"
tx = D["x"].query("device_id == @dev").sort_values("txn_time", ascending=False)
mt = D["m"].query("device_id == @dev")
info = D["d"].query("device_id == @dev").iloc[0]
fig = plt.figure(figsize=(16, 9), facecolor=BG)
fig.text(0.03, 0.95, f"Device Detail — {dev}   (страница drill-through, скрыта)", fontsize=18, weight="bold", color=NAVY)
sr = (tx.status == "success").mean()
for i, (l, v) in enumerate([("device_id", dev), ("model", info.model), ("location", info.location), ("Success Rate", f"{sr:.1%}")]):
    kpi(fig.add_axes([0.03 + i * 0.2425, 0.80, 0.225, 0.11]), l, v, NAVY)
    if i == 2: fig.axes[-1].texts[1].set_fontsize(13)
ax = fig.add_axes([0.04, 0.12, 0.50, 0.60]); ax.set_title("Транзакции устройства (последние 12)", loc="left", fontsize=11, color=NAVY)
tt = tx.head(12)[["txn_time", "txn_type", "amount", "status", "error_code"]].copy()
tt["txn_time"] = tt["txn_time"].dt.strftime("%d.%m %H:%M"); tt["amount"] = tt["amount"].astype(int); tt["error_code"] = tt["error_code"].fillna("")
table(ax, tt, fs=8.5, hl=lambda i: tt.iloc[i]["status"] != "success")
ax = fig.add_axes([0.60, 0.47, 0.36, 0.25], facecolor="white"); ec = tx[tx.status == "failed"]["error_code"].value_counts()
ax.bar(ec.index, ec.values, color=RED); ax.set_title("Коды ошибок (failed)", loc="left", fontsize=11, color=NAVY)
for s in ("top", "right"): ax.spines[s].set_visible(False)
ax = fig.add_axes([0.60, 0.12, 0.36, 0.28]); ax.set_title("Заявки устройства", loc="left", fontsize=11, color=NAVY)
if len(mt):
    q = mt[["ticket_id", "fault_time", "resolve_time", "fault_type"]].copy()
    q["MTTR, ч"] = ((mt.resolve_time - mt.fault_time).dt.total_seconds() / 3600).round(1)
    for c in ("fault_time", "resolve_time"): q[c] = q[c].dt.strftime("%d.%m %H:%M")
else:
    q = pd.DataFrame({"ticket_id": ["заявок нет"], "fault_time": [""], "resolve_time": [""], "fault_type": [""], "MTTR, ч": [""]})
table(ax, q, fs=8.5)
footer(fig); fig.savefig(S / "powerbi" / "detail.png", dpi=110); plt.close(fig)

# ------------------------------------------------- drill-through меню (схема)
fig, ax = plt.subplots(figsize=(8, 5), facecolor="white"); ax.axis("off"); ax.set_xlim(0, 10); ax.set_ylim(0, 6.5)
ax.text(0.2, 6.1, "Overview → таблица «Аномальные устройства»", fontsize=11, weight="bold", color=NAVY)
for i, (a, b, c) in enumerate([("ATM-007", "Model-X", "онлайн, но нет подключения > 3 дней"), ("ATM-014", "Model-Y", "…"), ("ATM-022", "Model-X", "…")]):
    ax.text(0.3, 5.1 - i * 0.55, f"{a}    {b}    {c}", fontsize=9, family="DejaVu Sans Mono", bbox=dict(fc="#EEF2F7" if i else "#DCE7F5", ec="none", pad=3))
ax.add_patch(plt.Rectangle((2.2, 1.0), 4.6, 2.4, fc="white", ec="#888", lw=1))
for j, (txt, bold) in enumerate([("Копировать", 0), ("Показать как таблицу", 0), ("Детализация  ▸", 1)]):
    ax.text(2.4, 3.0 - j * 0.5, txt, fontsize=10, weight="bold" if bold else "normal", color=NAVY if bold else "#333")
ax.add_patch(plt.Rectangle((6.9, 1.6), 2.9, 0.9, fc="white", ec="#888"))
ax.text(7.05, 2.05, "Device Detail", fontsize=10, weight="bold", color=NAVY)
ax.annotate("", xy=(6.9, 2.05), xytext=(6.1, 2.5), arrowprops=dict(arrowstyle="->", color=RED))
ax.text(0.3, 0.3, "ПКМ по строке → «Детализация» → «Device Detail»", fontsize=9, color=GREY)
footer(fig, "Схематичный рендер меню, не снимок экрана Power BI — заменить реальным скриншотом"); fig.savefig(S / "powerbi" / "drillthrough_menu.png", dpi=110); plt.close(fig)

# ------------------------------------------------------------- Power Query
dc = D["dc"].copy()
cols = ["device_id", "model", "status", "last_online", "txn_count", "last_txn_time", "anomaly_reason"]
def fmt(df):
    df = df[cols].copy()
    for c in ("last_online", "last_txn_time"): df[c] = df[c].dt.strftime("%d.%m %H:%M").fillna("null")
    return df
fig, ax = plt.subplots(figsize=(14, 4.6), facecolor="white")
ax.set_title("device_consistency: LEFT JOIN (device_status_clean + агрегат transactions_clean) — первые 10 строк", loc="left", fontsize=11, color=NAVY)
table(ax, fmt(dc.head(10)), fs=8.5, hl=lambda i: dc.head(10).iloc[i]["anomaly_reason"] != "норма")
footer(fig); fig.savefig(S / "power_query" / "pq_merge.png", dpi=110, bbox_inches="tight"); plt.close(fig)
fig, ax = plt.subplots(figsize=(14, 3.0), facecolor="white")
ax.set_title("device_consistency: строки с anomaly_reason ≠ «норма»", loc="left", fontsize=11, color=NAVY)
table(ax, fmt(dc.query("anomaly_reason != 'норма'")), fs=8.5, hl=lambda i: True)
footer(fig); fig.savefig(S / "power_query" / "pq_anomaly.png", dpi=110, bbox_inches="tight"); plt.close(fig)

# --------------------------------------------------------------------- SQL
def psql(sql):
    out = subprocess.run(["su", "postgres", "-c", f'psql -d terminal_analytics --csv -c "{sql}"'], capture_output=True, text=True).stdout
    return pd.read_csv(io.StringIO(out), keep_default_na=False)
q1 = psql("SELECT d.device_id, d.model, d.location, COUNT(*) AS total_txns, ROUND(100.0*COUNT(*) FILTER (WHERE t.status='success')/COUNT(*),1) AS success_rate_pct FROM devices d JOIN transactions t USING (device_id) GROUP BY 1,2,3 ORDER BY success_rate_pct, total_txns DESC LIMIT 8")
q2 = psql("SELECT d.device_id, d.model, d.last_online, MAX(t.txn_time) AS last_txn_time, COUNT(t.txn_id) AS txn_count FROM devices d LEFT JOIN transactions t USING (device_id) WHERE d.status='online' GROUP BY 1,2,3 HAVING MAX(t.txn_time) IS NULL OR MAX(t.txn_time) < TIMESTAMP '2026-09-30 23:59:59' - INTERVAL '3 days' ORDER BY last_txn_time NULLS FIRST")
for name, df, title in [("query_success_rate", q1, "Запрос 1. Коэффициент успешности по устройствам (8 худших)"),
                        ("query_anomaly_devices", q2, "Запрос 2. Устройства-«лжецы»: online, но нет транзакций > 3 дней")]:
    fig, ax = plt.subplots(figsize=(12, 1.2 + 0.36 * len(df)), facecolor="white")
    ax.set_title(title + "\nterminal_analytics=# результат (PostgreSQL 16)", loc="left", fontsize=10.5, color=NAVY)
    table(ax, df, fs=9); footer(fig, "Рендер результата реального запроса (PostgreSQL), не снимок экрана клиента — заменить скриншотом pgAdmin/DBeaver")
    fig.savefig(S / "sql" / f"{name}.png", dpi=110, bbox_inches="tight"); plt.close(fig)

# ------------------------------------------------------------------- Excel
from openpyxl import load_workbook
wbv = load_workbook(ROOT / "excel" / "terminal_analysis.xlsx", data_only=True)
wbf = load_workbook(ROOT / "excel" / "terminal_analysis.xlsx")
pv = [[c for c in row] for row in wbv["Pivot"].iter_rows(min_row=4, max_row=16, max_col=6, values_only=True)]
df = pd.DataFrame(pv[1:], columns=pv[0])
df["sum_amount, ₽"] = df["sum_amount, ₽"].map(lambda v: f"{v:,.0f}".replace(",", " ") if v is not None else "")
df["success_rate"] = df["success_rate"].map(lambda v: f"{v:.1%}" if isinstance(v, float) else v)
fig, ax = plt.subplots(figsize=(10, 4.6), facecolor="white"); ax.set_title("Лист Pivot: model × location (первые 12 строк)", loc="left", fontsize=11, color=NAVY)
table(ax, df, fs=9); footer(fig, "Рендер листа Pivot, не снимок экрана Excel — заменить скриншотом"); fig.savefig(S / "excel" / "pivot.png", dpi=110, bbox_inches="tight"); plt.close(fig)
rows = []
for r in range(2, 8):
    rows.append([wbv["Raw"].cell(r, 2).value, wbv["Raw"].cell(r, 8).value, wbf["Raw"].cell(r, 8).value, wbv["Raw"].cell(r, 11).value, wbf["Raw"].cell(r, 11).value])
df = pd.DataFrame(rows, columns=["device_id", "model", "формула H (ВПР)", "service_cost", "формула K (ВПР)"])
fig, ax = plt.subplots(figsize=(12, 3.2), facecolor="white"); ax.set_title("Лист Raw: ВПР по device_id → model и по model → стоимость; итог — СУММЕСЛИ на листе Lookup", loc="left", fontsize=10.5, color=NAVY)
table(ax, df, fs=9); footer(fig, "Рендер формул листа Raw, не снимок экрана Excel — заменить скриншотом"); fig.savefig(S / "excel" / "vlookup.png", dpi=110, bbox_inches="tight"); plt.close(fig)
print("ok")
