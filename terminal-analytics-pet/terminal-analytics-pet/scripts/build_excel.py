#!/usr/bin/env python3
"""Сборка excel/terminal_analysis.xlsx: Raw (+ВПР), Pivot (формулы), Lookup (model_cost + СУММЕСЛИ)."""
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from metrics import load

ROOT = Path(__file__).resolve().parent.parent
t, d, m, dc, am, log = load()
HDR = PatternFill("solid", fgColor="1F3A5F"); HF = Font(bold=True, color="FFFFFF", name="Arial")
BASE = Font(name="Arial", size=10); BLUE = Font(name="Arial", size=10, color="0000FF")
thin = Side(style="thin", color="BBBBBB"); BOX = Border(top=thin, bottom=thin, left=thin, right=thin)

def head(ws, row, col, names):
    for i, n in enumerate(names):
        c = ws.cell(row, col + i, n); c.fill, c.font = HDR, HF
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

wb = Workbook()
# ------------------------------------------------------------------ Lookup
lk = wb.active; lk.title = "Lookup"
lk["A1"] = "model_cost — стоимость обслуживания одной транзакции, ₽ (условные значения для учебного проекта)"; lk["A1"].font = Font(name="Arial", bold=True)
head(lk, 3, 1, ["model", "cost_per_txn_rub"])
for i, (mod, cost) in enumerate([("Model-X", 18), ("Model-Y", 14), ("Model-Z", 16)], 4):
    lk.cell(i, 1, mod).font = BASE; c = lk.cell(i, 2, cost); c.font = BLUE
head(lk, 3, 4, ["model", "txns", "service_cost_rub (СУММЕСЛИ)", "success_rate"])
n = len(t) + 1
for i, mod in enumerate(["Model-X", "Model-Y", "Model-Z"], 4):
    lk.cell(i, 4, f"=A{i}")
    lk.cell(i, 5, f"=COUNTIF(Raw!$H$2:$H${n},D{i})")
    lk.cell(i, 6, f"=SUMIF(Raw!$H$2:$H${n},D{i},Raw!$K$2:$K${n})")
    lk.cell(i, 7, f"=SUMIFS(Raw!$J$2:$J${n},Raw!$H$2:$H${n},D{i})/E{i}")
    lk.cell(i, 7).number_format = "0.0%"
lk["D7"], lk["E7"], lk["F7"], lk["G7"] = "ИТОГО", "=SUM(E4:E6)", "=SUM(F4:F6)", f"=SUM(Raw!$J$2:$J${n})/E7"
lk["G7"].number_format = "0.0%"
for c in "DEFG": lk[f"{c}7"].font = Font(name="Arial", bold=True)
lk["A9"] = "Проверка: ИТОГО txns должно равняться числу строк Raw"; lk["A9"].font = Font(name="Arial", italic=True)
lk["D9"] = f'=IF(E7=COUNTA(Raw!$A$2:$A${n}),"OK","ОШИБКА")'
for col, w in zip("ABCDEFG", [12, 18, 3, 12, 10, 28, 14]): lk.column_dimensions[col].width = w

# --------------------------------------------------------------------- Raw
raw = wb.create_sheet("Raw", 0)
cols = ["txn_id", "device_id", "txn_time", "txn_type", "amount", "status", "error_code"]
head(raw, 1, 1, cols + ["model (ВПР)", "location (ВПР)", "is_success", "service_cost (ВПР)"])
for r, row in enumerate(t[cols].itertuples(index=False), 2):
    for c, v in enumerate(row, 1):
        if v != v or v is None: v = None
        elif c == 3: v = v.to_pydatetime()
        cell = raw.cell(r, c, v); cell.font = BASE
        if c == 3: cell.number_format = "yyyy-mm-dd hh:mm:ss"
    raw.cell(r, 8, f"=VLOOKUP(B{r},$N$2:$S${len(d)+1},3,FALSE)").font = BASE
    raw.cell(r, 9, f"=VLOOKUP(B{r},$N$2:$S${len(d)+1},2,FALSE)").font = BASE
    raw.cell(r, 10, f'=IF(F{r}="success",1,0)').font = BASE
    raw.cell(r, 11, f"=VLOOKUP(H{r},Lookup!$A$4:$B$6,2,FALSE)").font = BASE
dcols = ["device_id", "location", "model", "last_online", "status", "cash_level"]
head(raw, 1, 14, dcols)
for r, row in enumerate(d[dcols].itertuples(index=False), 2):
    for c, v in enumerate(row, 14):
        if v != v: v = None
        elif c == 17: v = v.to_pydatetime()
        cell = raw.cell(r, c, v); cell.font = BASE
        if c == 17: cell.number_format = "yyyy-mm-dd hh:mm"
raw.freeze_panes = "A2"
for col, w in zip("ABCDEFGHIJK", [10, 10, 19, 15, 9, 9, 10, 11, 28, 10, 12]): raw.column_dimensions[col].width = w
for col, w in zip("NOPQRS", [10, 30, 10, 17, 12, 10]): raw.column_dimensions[col].width = w
raw.row_dimensions[1].height = 30

# ------------------------------------------------------------------- Pivot
pv = wb.create_sheet("Pivot", 1)
pv["A1"] = "Сводка: model × location (формулы COUNTIFS/SUMIFS по листу Raw)"; pv["A1"].font = Font(name="Arial", bold=True, size=12)
pv["A2"] = ("Вычисляемое поле Success Rate в сводной таблице Excel не умеет считать «долю status = success» — "
            "поэтому в Raw добавлен столбец is_success (1/0), а Success Rate = SUM(is_success) / COUNT.")
pv["A2"].font = Font(name="Arial", italic=True, size=9); pv["A2"].alignment = Alignment(wrap_text=True, vertical="top")
pv.merge_cells("A2:G2"); pv.row_dimensions[2].height = 38
head(pv, 4, 1, ["model", "location", "txns", "sum_amount, ₽", "success_txns", "success_rate"])
pairs = d.sort_values(["model", "location"])[["model", "location"]].values.tolist()
r0 = 5
for i, (mod, loc) in enumerate(pairs):
    r = r0 + i
    pv.cell(r, 1, mod); pv.cell(r, 2, loc)
    pv.cell(r, 3, f"=COUNTIFS(Raw!$H$2:$H${n},A{r},Raw!$I$2:$I${n},B{r})")
    pv.cell(r, 4, f"=SUMIFS(Raw!$E$2:$E${n},Raw!$H$2:$H${n},A{r},Raw!$I$2:$I${n},B{r})")
    pv.cell(r, 5, f"=SUMIFS(Raw!$J$2:$J${n},Raw!$H$2:$H${n},A{r},Raw!$I$2:$I${n},B{r})")
    pv.cell(r, 6, f'=IF(C{r}=0,"—",E{r}/C{r})'); pv.cell(r, 6).number_format = "0.0%"
    pv.cell(r, 4).number_format = "#,##0"
    for c in range(1, 7): pv.cell(r, c).font = BASE
rl = r0 + len(pairs)
pv.cell(rl, 1, "ИТОГО"); pv.cell(rl, 3, f"=SUM(C{r0}:C{rl-1})"); pv.cell(rl, 4, f"=SUM(D{r0}:D{rl-1})")
pv.cell(rl, 5, f"=SUM(E{r0}:E{rl-1})"); pv.cell(rl, 6, f"=E{rl}/C{rl}")
pv.cell(rl, 4).number_format = "#,##0"; pv.cell(rl, 6).number_format = "0.0%"
for c in range(1, 7): pv.cell(rl, c).font = Font(name="Arial", bold=True)
for col, w in zip("ABCDEF", [11, 32, 8, 15, 13, 13]): pv.column_dimensions[col].width = w
pv.freeze_panes = "A5"
wb.save(ROOT / "excel" / "terminal_analysis.xlsx")
print("saved")
